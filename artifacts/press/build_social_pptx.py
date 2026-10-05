#!/usr/bin/env python3
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Emu

ROOT = Path("/workspace/artifacts/press/social")
OUT = Path("/workspace/artifacts/DigitalGuardian-social.pptx")
PUBLIC = Path("/workspace/public/campaign/DigitalGuardian-social.pptx")

NOTES = [
    "Post 1 — hook. Caption: The black box for your real life. Stay protected. Preserve the truth. Not another SOS button.",
    "Post 2 — the gap. Caption: Late meeting. First date. Rideshare. If something happens, they don’t want guesses. They want the record.",
    "Post 3 — the system. Caption: Prevent. Detect. Preserve. Escalate. Reconstruct. One capsule. No panic button.",
    "Post 4 — the session. Caption: “If I don’t check in, start my protocol.” Guardian records facts. It does not decide danger.",
    "Post 5 — circle. Caption: Your people. Your rules. Missed check-ins follow only the protocol you wrote, including dead-man escrow.",
    "Post 6 — close. Caption: Start a Guardian Session. Write your protocol. Keep the record original. Observed events. Never a verdict.",
]

prs = Presentation()
prs.slide_width = Inches(7.5)
prs.slide_height = Inches(13.333333)
blank = prs.slide_layouts[6]

for i in range(1, 7):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(
        str(ROOT / f"c{i}.png"),
        Emu(0),
        Emu(0),
        width=prs.slide_width,
        height=prs.slide_height,
    )
    slide.notes_slide.notes_text_frame.text = NOTES[i - 1]

prs.core_properties.title = "GuardianOS — Social"
prs.core_properties.subject = "Digital guardian OS. Stay protected. Preserve the truth."
prs.core_properties.author = "GuardianOS"
PUBLIC.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
prs.save(PUBLIC)
print("wrote", OUT)
